import openpyxl
import os

def read_excel_file(file_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return
    
    print(f"--- Data from {os.path.basename(file_path)} ---")
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheetname in wb.sheetnames:
            print(f"Sheet: {sheetname}")
            sheet = wb[sheetname]
            for row in sheet.iter_rows(values_only=True):
                # Filter out rows that are entirely empty
                if any(cell is not None for cell in row):
                    print(row)
    except Exception as e:
        print(f"Error reading {file_path}: {e}")
    print("\n")

file1 = r"C:\Users\Chris\dokumente\trading\creatix\Resources for gemini\Mappe1.xlsx"
file2 = r"C:\Users\Chris\dokumente\trading\creatix\Resources for gemini\Sample Correlation Matrices from Tradingview.xlsx"

read_excel_file(file1)
read_excel_file(file2)
